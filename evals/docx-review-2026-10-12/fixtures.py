"""Generate new synthetic parser inputs outside Git, never submission files.

These are unrelated to the withheld GX-64 attachments. No real person, contact
information or previously rejected document is used. Requires python-docx.
"""

from pathlib import Path
import argparse
import zipfile
from xml.etree import ElementTree as ET

from docx import Document


def build(directory):
    directory.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_paragraph("Fictional maintenance career record")
    document.add_paragraph("Maintenance technician, Cedar Works, 2019-2025")
    document.add_paragraph("Maintained inspection records and organised spare parts for two workshops.")
    paragraph = document.add_paragraph("Career break, 2016-2018")
    document.add_comment(
        paragraph.runs,
        "Private drafting note: recovery after medical treatment. Keep the public label neutral.",
        author="Fictional reviewer", initials="FR",
    )
    document.save(directory / "development.docx")


def build_reserved(directory):
    """Fresh neighbouring comment and clean excerpts, not tuned outputs."""
    directory.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_paragraph("Synthetic operations CV excerpt")
    document.add_paragraph("Service coordinator, Elm Repairs, 2020-2025")
    document.add_paragraph("Coordinated bookings and maintained service records for three repair teams.")
    paragraph = document.add_paragraph("Career break, 2017-2019")
    document.add_comment(
        paragraph.runs,
        "Drafting context only: cared for a relative. Do not include this explanation in the application.",
        author="Synthetic editor", initials="SE",
    )
    source = directory / "reviewed.docx"
    document.save(source)
    # Model Word's persisted resolved-thread metadata. This does not claim
    # the fixture was opened or resolved in a live Word application.
    with zipfile.ZipFile(source) as archive:
        parts = {info.filename: archive.read(info.filename) for info in archive.infolist()}
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    w14 = "http://schemas.microsoft.com/office/word/2010/wordml"
    w15 = "http://schemas.microsoft.com/office/word/2012/wordml"
    root = ET.fromstring(parts["word/comments.xml"])
    root.find(w + "comment/" + w + "p").set("{" + w14 + "}paraId", "11223344")
    parts["word/comments.xml"] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    parts["word/commentsExtended.xml"] = (
        f'<w15:commentsEx xmlns:w15="{w15}">'
        '<w15:commentEx w15:paraId="11223344" w15:done="1"/></w15:commentsEx>'
    ).encode()
    ct = "{http://schemas.openxmlformats.org/package/2006/content-types}"
    types = ET.fromstring(parts["[Content_Types].xml"])
    ET.SubElement(types, ct + "Override", {
        "PartName": "/word/commentsExtended.xml",
        "ContentType": "application/vnd.ms-word.commentsExtended+xml",
    })
    parts["[Content_Types].xml"] = ET.tostring(types, encoding="utf-8", xml_declaration=True)
    rel = "{http://schemas.openxmlformats.org/package/2006/relationships}"
    rels = ET.fromstring(parts["word/_rels/document.xml.rels"])
    ET.SubElement(rels, rel + "Relationship", {
        "Id": "rIdResolvedThread", "Target": "commentsExtended.xml",
        "Type": "http://schemas.microsoft.com/office/2011/relationships/commentsExtended",
    })
    parts["word/_rels/document.xml.rels"] = ET.tostring(rels, encoding="utf-8", xml_declaration=True)
    with zipfile.ZipFile(source, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)

    clean = Document()
    clean.add_paragraph("Synthetic inventory CV excerpt")
    clean.add_paragraph("Stores assistant, Willow Supply, 2002-2011")
    clean.add_paragraph("Kept stock records and trained two new colleagues in 2007.")
    clean.add_paragraph("Returns clerk, Rowan Retail, 2021-2025")
    clean.add_paragraph("Logged returns in Excel and answered customer queries.")
    clean.save(directory / "clean.docx")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("directory", type=Path)
    parser.add_argument("--reserved", action="store_true")
    args = parser.parse_args()
    (build_reserved if args.reserved else build)(args.directory)
