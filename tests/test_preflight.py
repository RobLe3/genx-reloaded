"""Synthetic local-file checks. These do not test a real ATS."""

import hashlib
import importlib.util
import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "skills/genx-reloaded/scripts/cv_preflight.py"


class PreflightTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.work = Path(self.temp.name)

    def run_cli(self, source, *options):
        return subprocess.run(
            [sys.executable, str(SCRIPT), str(source), *map(str, options)],
            capture_output=True, text=True, check=False
        )

    def test_text_expected_phrases_and_source_preservation(self):
        source = self.work / "cv.txt"
        source.write_text(
            "Alex Example\nQuality Lead\n2010\u20132020\n"
            "Led a production team and maintained quality records for three lines.\n",
            encoding="utf-8",
        )
        before = hashlib.sha256(source.read_bytes()).hexdigest()
        result = self.run_cli(source, "--require", "2010-2020", "--require", "alex example")
        report = json.loads(result.stdout)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(report["status"], "no_local_flags")
        self.assertEqual(before, hashlib.sha256(source.read_bytes()).hexdigest())

    def test_missing_phrase_is_a_review_observation(self):
        source = self.work / "cv.md"
        source.write_text("A useful professional history with meaningful text. " * 4, encoding="utf-8")
        result = self.run_cli(source, "--require", "Never Supplied Employer")
        self.assertEqual(result.returncode, 0)
        report = json.loads(result.stdout)
        self.assertEqual(report["status"], "review_needed")
        self.assertIn("missing_expected_text", {f["code"] for f in report["findings"]})

    @unittest.skipUnless(importlib.util.find_spec("docx"), "python-docx not installed")
    def test_docx_flags_header_tables_hidden_and_tracked_changes(self):
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        doc = Document()
        doc.add_paragraph("Alex Example, Quality Lead")
        doc.add_paragraph("Led a production team and maintained operational quality records. " * 3)
        doc.sections[0].header.paragraphs[0].text = "Alex Example | alex@example.invalid"
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "A table with factual work experience."
        doc.add_paragraph().add_run("Hidden keyword").font.hidden = True
        insertion = OxmlElement("w:ins")
        insertion.set(qn("w:id"), "1")
        doc._element.body.append(insertion)
        source = self.work / "cv.docx"
        doc.save(source)
        report = json.loads(self.run_cli(source, "--require", "alex@example.invalid").stdout)
        codes = {f["code"] for f in report["findings"]}
        self.assertTrue({
            "docx_tables", "docx_hidden_text", "docx_tracked_changes",
            "header_footer_only_email", "header_footer_only_expected_text",
        }.issubset(codes))

    @unittest.skipUnless(
        importlib.util.find_spec("reportlab") and importlib.util.find_spec("pypdf"),
        "PDF test dependencies not installed",
    )
    def test_pdf_sparse_page_and_required_text(self):
        from reportlab.pdfgen import canvas
        source = self.work / "cv.pdf"
        pdf = canvas.Canvas(str(source))
        pdf.drawString(50, 750, "Alex Example, Quality Lead")
        pdf.drawString(50, 720, "Maintained quality records and supervised production teams.")
        pdf.showPage()
        pdf.rect(50, 650, 200, 80, fill=1)
        pdf.save()
        report = json.loads(self.run_cli(source, "--require", "Alex Example").stdout)
        self.assertEqual(report["structure"]["pages"], 2)
        self.assertTrue(report["expected_text"][0]["found_in_main_text"])
        self.assertIn("sparse_pdf_page", {f["code"] for f in report["findings"]})

    def test_json_and_text_outputs(self):
        source = self.work / "cv.txt"
        source.write_text("Supported professional experience and current capabilities. " * 4, encoding="utf-8")
        target = self.work / "report.json"
        extracted = self.work / "extracted.txt"
        result = self.run_cli(source, "--report", target, "--text-out", extracted)
        self.assertEqual(result.returncode, 0)
        self.assertEqual(json.loads(target.read_text()), json.loads(result.stdout))
        self.assertEqual(extracted.read_text(), source.read_text())

    def test_output_cannot_replace_source(self):
        source = self.work / "cv.txt"
        source.write_text("Preserve this source.", encoding="utf-8")
        result = self.run_cli(source, "--report", source)
        self.assertEqual(result.returncode, 2)
        self.assertEqual(source.read_text(), "Preserve this source.")

    def test_malformed_file_is_not_a_pass(self):
        source = self.work / "broken.docx"
        source.write_text("This is not a Word file.", encoding="utf-8")
        result = self.run_cli(source)
        self.assertEqual(result.returncode, 1)
        self.assertEqual(json.loads(result.stdout)["status"], "inspection_failed")

    def test_whitespace_requirement_is_rejected(self):
        source = self.work / "cv.txt"
        source.write_text("Factual work experience and clear practical evidence. " * 4, encoding="utf-8")
        result = self.run_cli(source, "--require", " ")
        self.assertEqual(result.returncode, 2)

    def test_phrase_does_not_match_inside_another_word(self):
        source = self.work / "cv.txt"
        source.write_text("Planning NoSQL migrations for multiple teams. " * 4, encoding="utf-8")
        report = json.loads(self.run_cli(source, "--require", "Ann", "--require", "SQL", "--require", "NoSQL").stdout)
        self.assertEqual([row["found_in_main_text"] for row in report["expected_text"]], [False, False, True])

    def test_hard_link_cannot_overwrite_source(self):
        source = self.work / "cv.txt"
        source.write_text("Original career record.", encoding="utf-8")
        alias = self.work / "report.json"
        os.link(source, alias)
        result = self.run_cli(source, "--report", alias)
        self.assertEqual(result.returncode, 2)
        self.assertEqual(source.read_text(), "Original career record.")

    def test_output_aliases_are_rejected_before_writing(self):
        source = self.work / "cv.txt"
        source.write_text("Supported work history. " * 6, encoding="utf-8")
        report = self.work / "report.json"
        report.write_text("Existing report", encoding="utf-8")
        alias = self.work / "extracted.txt"
        os.link(report, alias)
        result = self.run_cli(source, "--report", report, "--text-out", alias)
        self.assertEqual(result.returncode, 2)
        self.assertEqual(report.read_text(), "Existing report")

    @unittest.skipUnless(importlib.util.find_spec("docx"), "python-docx not installed")
    def test_hidden_and_deleted_text_cannot_satisfy_expected_text(self):
        from docx import Document
        from docx.oxml import OxmlElement
        doc = Document()
        doc.add_paragraph("A clear factual career history with enough visible body text. " * 3)
        doc.add_paragraph().add_run("FabricatedCredential").font.hidden = True
        doc.add_paragraph().add_run("VisibleSkill").font.hidden = False
        for tag, phrase in [("w:del", "DeletedClaim"), ("w:moveFrom", "MovedClaim")]:
            change = OxmlElement(tag)
            run = OxmlElement("w:r")
            text = OxmlElement("w:t")
            text.text = phrase
            run.append(text)
            change.append(run)
            doc.add_paragraph()._p.append(change)
        source = self.work / "cv.docx"
        doc.save(source)
        result = self.run_cli(source, "--require", "FabricatedCredential", "--require", "DeletedClaim", "--require", "MovedClaim", "--require", "VisibleSkill")
        report = json.loads(result.stdout)
        self.assertEqual(result.returncode, 0)
        self.assertEqual([row["found_in_main_text"] for row in report["expected_text"]], [False, False, False, True])

    @unittest.skipUnless(importlib.util.find_spec("docx"), "python-docx not installed")
    def test_style_hidden_setting_requires_review(self):
        from docx import Document
        from docx.enum.style import WD_STYLE_TYPE
        doc = Document()
        doc.add_paragraph("Factual professional history with clear visible evidence. " * 3)
        style = doc.styles.add_style("HiddenStyle", WD_STYLE_TYPE.CHARACTER)
        style.font.hidden = True
        doc.add_paragraph().add_run("Styled text", style="HiddenStyle")
        source = self.work / "cv.docx"
        doc.save(source)
        report = json.loads(self.run_cli(source).stdout)
        self.assertIn("docx_hidden_text", {f["code"] for f in report["findings"]})
        self.assertEqual(report["status"], "review_needed")

    @unittest.skipUnless(importlib.util.find_spec("docx"), "python-docx not installed")
    def test_explicit_column_definitions_are_reported(self):
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        doc = Document()
        doc.add_paragraph("A clear professional history with concrete information. " * 3)
        cols = doc.sections[0]._sectPr.find(qn("w:cols"))
        cols.attrib.pop(qn("w:num"), None)
        cols.append(OxmlElement("w:col"))
        cols.append(OxmlElement("w:col"))
        source = self.work / "cv.docx"
        doc.save(source)
        report = json.loads(self.run_cli(source).stdout)
        self.assertEqual(report["structure"]["configured_columns"], 2)
        self.assertIn("docx_columns", {f["code"] for f in report["findings"]})

    def test_directory_is_an_inspection_failure(self):
        result = self.run_cli(self.work)
        self.assertEqual(result.returncode, 1)
        self.assertEqual(json.loads(result.stdout)["status"], "inspection_failed")


if __name__ == "__main__":
    unittest.main()
